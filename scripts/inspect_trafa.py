from __future__ import annotations

import json
from pathlib import Path
from urllib.request import Request, urlopen
from io import BytesIO

from openpyxl import load_workbook

SOURCES = {
    "vehicles": "https://www.trafa.se/globalassets/statistik/vagtrafik/fordon/2026/fordon-i-lan-och-kommuner-2025.xlsx",
    "injuries": "https://prod.trafa.se/globalassets/statistik/vagtrafik/vagtrafikskador/2025/vagtrafikskador-2025.xlsx",
    "service": "https://www.trafa.se/globalassets/statistik/kollektivtrafik/fardtjanst/2025/fardtjanst-och-riksfardtjanst-2025.xlsx",
}

MUNICIPALITIES = [
    "Arjeplog","Arvidsjaur","Boden","Gällivare","Haparanda","Jokkmokk","Kalix",
    "Kiruna","Luleå","Pajala","Piteå","Älvsbyn","Överkalix","Övertorneå",
]

def download(url: str) -> bytes:
    req = Request(url, headers={"User-Agent": "pendling-transport-lulea/1.0"})
    with urlopen(req, timeout=120) as r:
        return r.read()

def row_values(ws, row_idx: int):
    vals=[]
    for cell in ws[row_idx]:
        v=cell.value
        if isinstance(v, (int,float)) or v is None:
            vals.append(v)
        else:
            vals.append(str(v).strip())
    while vals and vals[-1] is None:
        vals.pop()
    return vals

def main():
    out={}
    for key,url in SOURCES.items():
        raw=download(url)
        wb=load_workbook(BytesIO(raw), data_only=True, read_only=True)
        info={"source":url,"sheets":[]}
        for ws in wb.worksheets:
            matches=[]
            max_scan=min(ws.max_row or 0, 1200)
            for i in range(1,max_scan+1):
                vals=row_values(ws,i)
                text=" | ".join("" if v is None else str(v) for v in vals)
                if any(name.casefold() in text.casefold() for name in MUNICIPALITIES):
                    matches.append({"row":i,"values":vals[:40]})
                    if len(matches)>=30:
                        break
            info["sheets"].append({
                "title":ws.title,
                "max_row":ws.max_row,
                "max_column":ws.max_column,
                "matches":matches,
            })
        out[key]=info

    path=Path("data/trafa_inspect.json")
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    print(json.dumps({k:[s["title"] for s in v["sheets"]] for k,v in out.items()},ensure_ascii=False,indent=2))

if __name__=="__main__":
    main()
