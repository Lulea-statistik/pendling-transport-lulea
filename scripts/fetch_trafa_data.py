from __future__ import annotations

import json
import re
from pathlib import Path
from urllib.request import Request, urlopen
from io import BytesIO

from openpyxl import load_workbook

SOURCES = {
    "vehicles": {
        "url": "https://www.trafa.se/globalassets/statistik/vagtrafik/fordon/2026/fordon-i-lan-och-kommuner-2025.xlsx",
        "year": 2025,
    },
    "injuries": {
        "url": "https://prod.trafa.se/globalassets/statistik/vagtrafik/vagtrafikskador/2025/vagtrafikskador-2025.xlsx",
        "year": 2025,
    },
    "service": {
        "url": "https://www.trafa.se/globalassets/statistik/kollektivtrafik/fardtjanst/2025/fardtjanst-och-riksfardtjanst-2025.xlsx",
        "year": 2025,
    },
}

MUNICIPALITIES = [
    "Arjeplog","Arvidsjaur","Boden","Gällivare","Haparanda","Jokkmokk","Kalix",
    "Kiruna","Luleå","Pajala","Piteå","Älvsbyn","Överkalix","Övertorneå",
]
NAME_BY_NORM = {x.casefold(): x for x in MUNICIPALITIES}

def download(url: str) -> bytes:
    req = Request(url, headers={"User-Agent": "pendling-transport-lulea/1.0"})
    with urlopen(req, timeout=180) as r:
        return r.read()

def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        v=value.strip()
        if v in {"–","-","..","."}:
            return None
        return v
    return value

def normalize_name(value):
    text=str(value or "").strip().casefold()
    return NAME_BY_NORM.get(text)

def useful_rows(ws, max_cols=60):
    max_col=min(max_cols, ws.max_column or max_cols)
    for idx,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row or 0,1200),min_col=1,max_col=max_col,values_only=True),start=1):
        yield idx,[clean(v) for v in row]

def combined_headers(header_rows, col_count):
    labels=[]
    for c in range(col_count):
        parts=[]
        for row in header_rows:
            if c>=len(row):
                continue
            v=row[c]
            if v in ("",None):
                continue
            text=re.sub(r"\s+"," ",str(v)).strip()
            if not text or text.isdigit():
                continue
            if text not in parts:
                parts.append(text)
        if parts:
            label=" · ".join(parts[-3:])
        else:
            label=f"Kolumn {c+1}"
        labels.append(label[:180])
    return labels

def extract_workbook(key, cfg):
    raw=download(cfg["url"])
    wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)
    result={
        "source":cfg["url"],
        "year":cfg["year"],
        "sheets":[],
    }

    for ws in wb.worksheets:
        scanned=list(useful_rows(ws))
        matches=[]
        for row_idx,vals in scanned:
            found=None
            found_col=None
            for c,v in enumerate(vals):
                name=normalize_name(v)
                if name:
                    found=name
                    found_col=c
                    break
            if found:
                matches.append((row_idx,found,found_col,vals))

        if not matches:
            continue

        first_row=matches[0][0]
        header_rows=[
            vals for idx,vals in scanned
            if max(1,first_row-12)<=idx<first_row
            and any(v not in ("",None) for v in vals)
        ]
        max_used=max(
            max((i+1 for i,v in enumerate(vals) if v not in ("",None)),default=0)
            for _,_,_,vals in matches
        )
        labels=combined_headers(header_rows,max_used)

        rows=[]
        seen=set()
        for _,name,_,vals in matches:
            if name in seen:
                continue
            seen.add(name)
            data={}
            for i in range(min(max_used,len(vals))):
                value=vals[i]
                if value in ("",None):
                    continue
                label=labels[i]
                # Skip municipality name/code columns from metric list.
                if normalize_name(value)==name:
                    continue
                if i<=1 and (isinstance(value,(int,float)) or str(value).isdigit()):
                    continue
                data[label]=value
            rows.append({"municipality":name,"values":data})

        if rows:
            result["sheets"].append({
                "title":ws.title,
                "columns":labels,
                "rows":rows,
            })

    return result

def main():
    out={}
    for key,cfg in SOURCES.items():
        print("Hämtar",key,cfg["url"])
        out[key]=extract_workbook(key,cfg)
        print(key,"kommunblad:",[s["title"] for s in out[key]["sheets"]])

    for target in [Path("data/trafa.json"),Path("docs/data/trafa.json")]:
        target.parent.mkdir(parents=True,exist_ok=True)
        target.write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding="utf-8")

    print("Klart", {k:len(v["sheets"]) for k,v in out.items()})

if __name__=="__main__":
    main()
