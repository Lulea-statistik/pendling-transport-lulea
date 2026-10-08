from __future__ import annotations

import json
from pathlib import Path
from urllib.request import Request, urlopen
from io import BytesIO
from openpyxl import load_workbook

SOURCES = {
    "road": "https://www.nvdb.se/globalassets/upload/anvandare/statistik/antal-kilometer-bilnat-per-vaghallare-efter-lan-och-kommun.xlsx",
    "cycle": "https://www.nvdb.se/globalassets/upload/anvandare/statistik/antal-meter-cykelnat-per-vaghallare-efter-lan-och-kommun.xlsx",
}

MUNICIPALITIES = [
    "Arjeplog","Arvidsjaur","Boden","Gällivare","Haparanda","Jokkmokk","Kalix",
    "Kiruna","Luleå","Pajala","Piteå","Älvsbyn","Överkalix","Övertorneå",
]

def download(url):
    req=Request(url,headers={"User-Agent":"pendling-transport-lulea/1.0"})
    with urlopen(req,timeout=120) as r:
        return r.read()

def clean(v):
    if v is None:
        return ""
    if isinstance(v,str):
        return v.strip()
    return v

def main():
    out={}
    for key,url in SOURCES.items():
        raw=download(url)
        wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)
        info={"url":url,"sheets":[]}
        for ws in wb.worksheets:
            rows=[]
            for i,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row or 0,1200),values_only=True),start=1):
                vals=[clean(v) for v in row]
                text=" | ".join(str(v) for v in vals if v not in ("",None))
                if any(m.casefold() in text.casefold() for m in MUNICIPALITIES):
                    rows.append({"row":i,"values":vals[:40]})
            top=[]
            for i,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row or 0,50),values_only=True),start=1):
                vals=[clean(v) for v in row]
                if any(v not in ("",None) for v in vals):
                    top.append({"row":i,"values":vals[:40]})
            info["sheets"].append({
                "title":ws.title,
                "max_row":ws.max_row,
                "max_column":ws.max_column,
                "top":top,
                "matches":rows[:50],
            })
        out[key]=info

    p=Path("data/nvdb_inspect.json")
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    print(json.dumps({k:[s["title"] for s in v["sheets"]] for k,v in out.items()},ensure_ascii=False,indent=2))

if __name__=="__main__":
    main()
