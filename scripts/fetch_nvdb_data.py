from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

SOURCES = {
    "road": {
        "url": "https://www.nvdb.se/globalassets/upload/anvandare/statistik/antal-kilometer-bilnat-per-vaghallare-efter-lan-och-kommun.xlsx",
        "unit": "km",
        "label": "Bilvägnät",
    },
    "cycle": {
        "url": "https://www.nvdb.se/globalassets/upload/anvandare/statistik/antal-meter-cykelnat-per-vaghallare-efter-lan-och-kommun.xlsx",
        "unit": "m",
        "label": "Cykelvägnät",
    },
}

NBR = {
    "arjeplog","arvidsjaur","boden","gällivare","haparanda","jokkmokk","kalix",
    "kiruna","luleå","pajala","piteå","älvsbyn","överkalix","övertorneå"
}

def download(url: str) -> bytes:
    req = Request(url, headers={"User-Agent": "pendling-transport-lulea/1.0"})
    with urlopen(req, timeout=180) as r:
        return r.read()

def to_num(v):
    if v is None or v == "" or str(v).strip() in {"-","–",".."}:
        return None
    try:
        return float(v)
    except Exception:
        return None

def extract(source):
    raw = download(source["url"])
    wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        title = str(ws.title).strip()
        if not title.isdigit():
            continue
        year = int(title)
        for row in ws.iter_rows(min_row=2, values_only=True):
            county = str(row[0] or "").strip()
            municipality = str(row[1] or "").strip()
            if county.casefold() != "norrbotten" or municipality.casefold() not in NBR:
                continue
            enskild = to_num(row[3] if len(row) > 3 else None)
            kommunal = to_num(row[4] if len(row) > 4 else None)
            statlig = to_num(row[5] if len(row) > 5 else None)
            rows.append({
                "year": year,
                "municipality": municipality,
                "private": enskild,
                "municipal": kommunal,
                "state": statlig,
                "total": sum(v for v in (enskild, kommunal, statlig) if v is not None),
            })
    rows.sort(key=lambda r: (r["year"], r["municipality"]))
    return {
        "source": source["url"],
        "unit": source["unit"],
        "label": source["label"],
        "years": sorted({r["year"] for r in rows}),
        "rows": rows,
    }

def main():
    out = {k: extract(v) for k,v in SOURCES.items()}
    for path in [Path("data/nvdb.json"), Path("docs/data/nvdb.json")]:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print({k: {"years": v["years"], "rows": len(v["rows"])} for k,v in out.items()})

if __name__ == "__main__":
    main()
